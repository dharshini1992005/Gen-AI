from docx import Document as DocxDocument
from langchain_core.documents import Document


def load_docx(file_path, paragraphs_per_section=25):
    """Read a .docx file and return Documents grouped into sections.
    Word files have no fixed pages, so each group of paragraphs is a 'section'."""
    doc = DocxDocument(file_path)

    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Include table content too
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    documents = []
    for i in range(0, len(lines), paragraphs_per_section):
        text = "\n".join(lines[i:i + paragraphs_per_section])
        documents.append(
            Document(
                page_content=text,
                metadata={
                    "source": file_path,
                    "page": i // paragraphs_per_section,
                    "format": "docx",
                },
            )
        )
    return documents